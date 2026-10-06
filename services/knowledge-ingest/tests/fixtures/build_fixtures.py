# SPDX-License-Identifier: Apache-2.0
# Copyright 2026 chanl-ai
"""Writes the fixture files next to this script. Run with `uv run python tests/fixtures/build_fixtures.py`.

The files are checked in; this script records how they were made so they can be rebuilt.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Flowable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HERE = Path(__file__).parent


def policy_pdf() -> None:
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(HERE / "policy.pdf"), pagesize=A4, title="Card Fees Policy")
    table = Table(
        [
            ["Card", "Annual fee", "Foreign transaction fee"],
            ["Classic", "$0", "3%"],
            ["Gold", "$95", "2%"],
            ["Platinum", "$450", "0%"],
        ]
    )
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    story: list[Flowable] = [
        Paragraph("Card Fees Policy", styles["Title"]),
        Paragraph("1 Annual fees", styles["Heading1"]),
        Paragraph(
            "Annual fees are charged on the account opening date and on each anniversary. "
            "A fee is refunded in full when the card is closed within 30 days of the charge.",
            styles["BodyText"],
        ),
        Spacer(1, 12),
        Paragraph("Table 1 Annual fees by card", styles["BodyText"]),
        table,
        PageBreak(),
        Paragraph("2 Late payment", styles["Heading1"]),
        Paragraph(
            "A late payment fee of $25 applies when the minimum payment is not received by the due date. "
            "The fee is waived once in any twelve-month period on request.",
            styles["BodyText"],
        ),
    ]
    doc.build(story)


def scanned_pdf() -> None:
    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=56)
    draw.text((100, 200), "LATE PAYMENT FEE", fill="black", font=font)
    draw.text((100, 300), "The fee is 25 dollars.", fill="black", font=font)
    image.save(HERE / "scanned.pdf", "PDF", resolution=150.0)


def procedures_docx() -> None:
    doc = Document()
    doc.core_properties.title = "Complaint Handling Procedure"
    doc.add_heading("Complaint Handling Procedure", level=0)
    doc.add_heading("1 Receiving a complaint", level=1)
    doc.add_paragraph("Log every complaint in the case system on the day it arrives.")
    doc.add_heading("1.1 Acknowledgement", level=2)
    doc.add_paragraph("Send an acknowledgement within two business days.")
    table = doc.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate(
        [("Channel", "Acknowledge within"), ("Email", "2 business days"), ("Letter", "5 business days")]
    ):
        table.cell(r, 0).text = a
        table.cell(r, 1).text = b
    doc.add_heading("2 Escalation", level=1)
    doc.add_paragraph("Escalate to the complaints team when the customer asks for it.")
    doc.save(str(HERE / "procedures.docx"))


def fees_xlsx() -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Fees"
    ws.append(["Product", "Balance band", "Monthly fee", "Internal note"])
    ws.append(["Everyday Account", "<= $250,000", 5.0, "legacy pricing"])
    ws.append(["Everyday Account", "$250,001 - $1,000,000", 0.0, "waived"])
    ws.append(["Premier Account", "over $1,000,000", 0.0, "relationship"])
    ws.append(["Saver Account", "any balance", 2.5, "review 2027"])
    wb.properties.title = "Account fee schedule"
    wb.save(HERE / "fees.xlsx")


HTML = """<!doctype html>
<html><head><title>Replace a lost card | Help centre</title>
<script>window.analytics = {track: function(){}};</script>
<style>.nav{color:red}</style></head>
<body>
<header class="site-header"><a href="/">Home</a> <a href="/help">Help</a></header>
<nav class="main-nav"><ul><li>Accounts</li><li>Cards</li><li>Loans</li></ul></nav>
<div id="cookie-banner">We use cookies to improve your experience. Accept all cookies</div>
<main>
<article>
<h1>Replace a lost card</h1>
<p class="lead">Report a lost card straight away so nobody else can use it.</p>
<h2>Report it</h2>
<p>Call the card line or freeze the card in the app. A replacement arrives in five business days.</p>
<h2>Fees</h2>
<table><caption>Replacement fees</caption>
<tr><th>Card</th><th>Fee</th></tr>
<tr><td>Classic</td><td>$0</td></tr>
<tr><td>Gold</td><td>$10</td></tr>
</table>
<div class="share-buttons">Share on social media</div>
</article>
</main>
<aside class="sidebar">Related: Travel insurance</aside>
<footer class="site-footer">Copyright 2026 Northfield Bank. All rights reserved.</footer>
</body></html>
"""


def help_html() -> None:
    (HERE / "help.html").write_text(HTML, encoding="utf-8")


if __name__ == "__main__":
    policy_pdf()
    scanned_pdf()
    procedures_docx()
    fees_xlsx()
    help_html()
