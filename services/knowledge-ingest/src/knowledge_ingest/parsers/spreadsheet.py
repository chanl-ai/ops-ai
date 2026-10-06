# SPDX-License-Identifier: Apache-2.0
# Copyright 2026 chanl-ai
"""XLSX workbooks and CSV files. Each sheet becomes a section holding one table.

The first non-empty row is the header row. Cell values are kept as the source shows them: a whole
number stored as a float is written without `.0`, and dates as ISO strings.
"""

from __future__ import annotations

import csv
import datetime as dt
import io

from openpyxl import load_workbook

from knowledge_ingest.models import ParsedDocument, ParsingSettings, SourceMetadata
from knowledge_ingest.parsers.builder import DocumentBuilder
from knowledge_ingest.parsers.text import decode

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_ROWS_PER_SHEET = 50_000


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value).strip()


def _add_sheet(builder: DocumentBuilder, name: str, rows: list[list[str]], keep: bool) -> None:
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    if len(rows) > MAX_ROWS_PER_SHEET + 1:
        builder.warnings.append(
            f"Sheet {name} has {len(rows) - 1} rows; only the first {MAX_ROWS_PER_SHEET} were read."
        )
        rows = rows[: MAX_ROWS_PER_SHEET + 1]
    builder.heading(name, 1)
    builder.table(rows[0], rows[1:], sheet=name, keep_as_table=keep)


def parse_xlsx(
    data: bytes,
    *,
    filename: str,
    parsing: ParsingSettings | None = None,
    fields: dict[str, str] | None = None,
) -> ParsedDocument:
    parsing = parsing or ParsingSettings()
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    builder = DocumentBuilder()
    try:
        for sheet in workbook.worksheets:
            rows = [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
            _add_sheet(builder, sheet.title, rows, parsing.tables)
        props = workbook.properties
        title = props.title or None
        source = SourceMetadata(
            filename=filename,
            mime_type=XLSX_MIME,
            size_bytes=len(data),
            title=title,
            author=props.creator or None,
            modified=props.modified.isoformat() if props.modified else None,
            fields=fields or {},
        )
    finally:
        workbook.close()
    return builder.build(title=title, source=source)


def parse_csv(
    data: bytes,
    *,
    filename: str,
    parsing: ParsingSettings | None = None,
    fields: dict[str, str] | None = None,
) -> ParsedDocument:
    parsing = parsing or ParsingSettings()
    text = decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    builder = DocumentBuilder()
    _add_sheet(builder, filename.rsplit(".", 1)[0], rows, parsing.tables)
    source = SourceMetadata(filename=filename, mime_type="text/csv", size_bytes=len(data), fields=fields or {})
    return builder.build(title=None, source=source)
